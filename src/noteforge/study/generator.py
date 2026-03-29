import os
import re
from typing import List, Dict, Any, Optional
from ..models.manager import ModelManager

class StudyGenerator:
    """Handles parsing and AI summarization for study guides."""
    def __init__(self):
        self._bart_pipeline = None

    def process_file(self, filepath: str, progress_callback=None) -> Dict[str, Any]:
        """Parses a file and generates notes."""
        if not os.path.exists(filepath):
            return {"error": f"File not found: {filepath}"}
            
        ext = os.path.splitext(filepath)[1].lower()
        try:
            if ext == ".txt":
                text = self._read_txt(filepath)
            elif ext == ".pptx":
                text = self._read_pptx(filepath)
            else:
                return {"error": f"Unsupported file type: {ext}"}
                
            if progress_callback:
                progress_callback("Analyzing content...", 0.2)
                
            notes = self.generate_notes(text, progress_callback)
            return {"notes": notes, "original_text": text}
        except Exception as e:
            return {"error": str(e)}

    def _read_txt(self, filepath: str) -> str:
        for enc in ['utf-8', 'latin-1', 'cp1252']:
            try:
                with open(filepath, "r", encoding=enc) as f:
                    return f.read()
            except:
                continue
        raise IOError("Could not read text file with supported encodings.")

    def _read_pptx(self, filepath: str) -> str:
        from pptx import Presentation
        prs = Presentation(filepath)
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
        return "\n".join(text_runs)

    def generate_notes(self, text: str, progress_callback=None) -> str:
        """Splits text into topics and summarizes each."""
        if not self._bart_pipeline:
            if progress_callback:
                progress_callback("Loading AI models...", 0.3)
            self._bart_pipeline = ModelManager.get_bart_pipeline()

        # Simple semantic splitting (by double newline or length)
        chunks = [c.strip() for c in text.split("\n\n") if len(c.strip()) > 100]
        if not chunks:
            chunks = [text[i:i+2000] for i in range(0, len(text), 2000)]

        results = []
        for i, chunk in enumerate(chunks):
            if progress_callback:
                progress_callback(f"Summarizing section {i+1}/{len(chunks)}...", 0.4 + (0.5 * i / len(chunks)))
            
            # Use AI to summarize
            summary = self._bart_pipeline(chunk, max_length=150, min_length=40, do_sample=False)[0]['summary_text']
            results.append(self._format_section(i + 1, chunk[:50] + "...", summary))
            
        return "\n\n".join(results)

    def _format_section(self, index: int, topic: str, summary: str) -> str:
        return f"### {index}. {topic.upper()}\n- **Overview:** {summary}"
